"""Build the SashaInfinity Education OS product document."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "SashaInfinity_Education_OS_Product_Document.docx"

NAVY = "18314F"
PALE_BLUE = "EEF4F8"
PALE_GRAY = "F7F8FA"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 98, 108)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def suppress_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    p_pr.append(borders)


def set_font(run, name="Arial", size=None, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = color


def add_paragraph(doc, text="", *, bold_lead=None, space_after=6, keep=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.13
    paragraph.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, size=10.5, bold=True)
        body = paragraph.add_run(text[len(bold_lead):])
        set_font(body, size=10.5)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=10.5)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08
        set_font(paragraph.add_run(item), size=10.5)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.08
        set_font(paragraph.add_run(item), size=10.5)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        set_font(run, size=16 if level == 1 else 12, bold=True, color=BLACK)
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run(str(header)), size=9.2, bold=True, color=RGBColor(255, 255, 255))
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, PALE_BLUE if row_index % 2 else "FFFFFF")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            set_font(paragraph.add_run(str(value)), size=8.8)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    prefix = paragraph.add_run("SashaInfinity Education OS   |   Product Document   |   ")
    set_font(prefix, size=8.5, color=GRAY)
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_end)
    set_font(run, size=8.5, color=GRAY)


def page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    add_footer(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = BLACK

    doc.core_properties.title = "SashaInfinity Education OS Product Document"
    doc.core_properties.subject = "Product architecture business verticals and production deployment"
    doc.core_properties.author = "SashaInfinity"
    doc.core_properties.keywords = "LMS, education OS, Meiporul, Seyappaduporul, Utporul"

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(16)
    suppress_paragraph_borders(title)
    set_font(title.add_run("SashaInfinity Education OS Product Document"), size=28, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    set_font(subtitle.add_run("Production architecture business verticals and deployment"), size=14, color=GRAY)
    add_paragraph(
        doc,
        "This document defines the complete operating model for SashaInfinity as one education platform with three independent business pillars. It records what has been built, how the pillars share identity and administration, how revenue remains attributable, and what must be completed in a real production environment before public launch.",
        space_after=18,
    )
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(42)
    meta.paragraph_format.space_after = Pt(6)
    set_font(meta.add_run("Release status"), size=9, bold=True, color=GRAY)
    add_paragraph(doc, "Production candidate validated 13 September 2026", space_after=12)
    meta2 = doc.add_paragraph()
    meta2.paragraph_format.space_after = Pt(6)
    set_font(meta2.add_run("Audience"), size=9, bold=True, color=GRAY)
    add_paragraph(doc, "Founders product leadership engineering operations finance and deployment partners")

    page_break(doc)

    add_heading(doc, "Product Summary")
    add_paragraph(
        doc,
        "SashaInfinity is organized as a shared platform kernel surrounded by three independently marketable product pillars. This solves the central structural problem: immersive learning, institution operations, and skill development no longer compete inside one undifferentiated LMS menu, while the business still maintains one account, one customer boundary, one administration plane, and one financial reporting model.",
    )
    add_paragraph(
        doc,
        "The main LMS remains the common discovery, authentication, checkout, learner, instructor, and support surface. Each pillar receives a dedicated public hostname and product experience. Customers may activate one or more pillars through entitlements without duplicate accounts or duplicated curriculum records.",
    )
    add_table(
        doc,
        ["Pillar", "Primary product", "Primary customers", "Revenue families"],
        [
            ["Meiporul", "3D AR VR curriculum and immersive lab deployment", "Learners schools colleges museums and lab partners", "Courses asset licenses subscriptions deployment milestones and support"],
            ["Seyappaduporul", "Tutoring and institution operations", "Schools colleges tutors parents and students", "Tuition institution plans books papers franchise and service fees"],
            ["Utporul", "Skill creation assessment credentials and careers", "Learners instructors employers and training providers", "Courses assessments credentials creator commerce and career services"],
        ],
        [1.05, 2.05, 1.75, 2.15],
    )
    add_heading(doc, "Core Product Decision", level=2)
    add_paragraph(
        doc,
        "Each source transaction stays authoritative. Course orders, institution fee ledgers, internship vouchers, commercial invoices, refunds, and payouts are not collapsed into a single editable table. The Sasha Control Center reads and reconciles those sources into a consolidated portfolio view. This protects accounting integrity while allowing each pillar to use a different commercial model.",
    )
    add_heading(doc, "Experience Principles", level=2)
    add_bullets(
        doc,
        [
            "A public visitor immediately understands which pillar they are exploring.",
            "A learner signs in once and sees only the capabilities and content granted to them.",
            "An instructor authors reusable learning components without having to operate campus administration.",
            "A customer organization is isolated by tenant and can activate capabilities across pillars.",
            "A platform administrator can diagnose user issues and operate every pillar from one audited console.",
        ],
    )

    page_break(doc)

    add_heading(doc, "Platform Architecture")
    add_paragraph(
        doc,
        "The architecture is a modular monolith for the core business domain with separate workers and services where isolation or resource control is required. React and Vite provide the web application, FastAPI owns APIs and business services, PostgreSQL is the system of record, Redis supports transient coordination, nginx is the only public ingress, and Docker Compose defines the deployable topology.",
    )
    add_table(
        doc,
        ["Layer", "Responsibilities", "Key implementation"],
        [
            ["Public experience", "Pillar discovery catalog checkout and learner access", "React 18 Vite Tailwind hostname aware routes"],
            ["Control plane", "Admin reporting tenants entitlements content people and money operations", "Sasha Control Center and role scoped workspaces"],
            ["Business services", "Learning institution commerce assessment and fulfillment rules", "FastAPI routers with service layer boundaries"],
            ["Platform kernel", "Identity tenancy domains entitlements audit outbox and finance", "SQLAlchemy models PostgreSQL Alembic migrations"],
            ["Isolated services", "Video extraction live meetings code execution and background work", "Streaming service Jitsi Judge0 private workers Redis"],
            ["Delivery edge", "TLS routing rate limits static media and subdomains", "nginx CDN object storage and DNS"],
        ],
        [1.25, 3.0, 2.75],
    )
    add_heading(doc, "Tenant Boundary", level=2)
    add_paragraph(
        doc,
        "PlatformTenant is the customer boundary. Memberships define who may operate inside it; domains map branded hostnames; entitlements activate features per pillar; audit events record administrative decisions; and outbox events make significant changes available to downstream integrations. Institution provisioning synchronizes into this boundary rather than creating another LMS.",
    )
    add_heading(doc, "Content Composition", level=2)
    add_paragraph(
        doc,
        "Course authoring is shared because quizzes, assignments, videos, games, live sessions, 3D objects, labs, coding tasks, certificates, and outcomes are reusable learning primitives. Every course still carries one canonical business vertical so public discovery, inventory, revenue, and performance reports remain unambiguous.",
    )

    page_break(doc)

    add_heading(doc, "Meiporul Immersive Learning")
    add_paragraph(
        doc,
        "Meiporul is the immersive curriculum and experiential learning business. It combines public discovery of AR and 3D assets with curriculum-bound activities, virtual labs, evidence collection, and physical lab deployment operations. It is not a campus-management module and it is not only an object gallery.",
    )
    add_heading(doc, "Learner and Instructor Capabilities", level=2)
    add_bullets(
        doc,
        [
            "Curated immersive courses with reusable GLB and 3D model assets.",
            "AR presentation and browser based 3D viewing tied to lessons and tasks.",
            "Virtual lab catalog with native guided and imported experiences.",
            "Lab Studio publication controls investigation notebooks and learner evidence.",
            "GeoGebra interactive content with explicit validation and curriculum placement.",
            "Scorable 3D match and verify tasks mastery evidence and learning outcomes.",
        ],
    )
    add_heading(doc, "Deployment Operations", level=2)
    add_table(
        doc,
        ["Operational area", "Capability", "Business value"],
        [
            ["Sites", "Customer location ownership and deployment state", "Separates physical delivery from course authoring"],
            ["Devices", "Fleet identity condition assignment and service state", "Supports asset accountability and replacement planning"],
            ["Safety", "Inspection records readiness gates and go live control", "Prevents unsafe or incomplete launch"],
            ["Milestones", "Rollout deliverables acceptance and revenue stages", "Supports milestone billing and delivery reporting"],
            ["Support", "Service tickets priority SLA state and resolution", "Enables AMC and field support revenue"],
        ],
        [1.15, 3.2, 2.65],
    )
    add_heading(doc, "Meiporul Commercial Model", level=2)
    add_paragraph(
        doc,
        "Revenue can be generated through immersive course sales, individual or institutional asset licensing, experience subscriptions, physical lab deployment milestones, and annual maintenance or field support. These offers are configured through the shared commercial catalog so Meiporul remains financially attributable without receiving a separate accounting stack.",
    )

    page_break(doc)

    add_heading(doc, "Seyappaduporul Tutoring and Institutions")
    add_paragraph(
        doc,
        "Seyappaduporul is the tutoring and school operations business. It owns the tools used to deliver recurring instruction and operate an institution. Campus functionality is a tenant scoped product workspace under SashaInfinity, not a hidden second platform inside the course system.",
    )
    add_heading(doc, "Teaching and Learning", level=2)
    add_bullets(
        doc,
        [
            "Live Jitsi classes scheduling join controls attendance recordings and reports.",
            "Tutor notebooks private resources note sharing and ebook distribution.",
            "Question banks practice paper generation priced paper access exams and results.",
            "Assignments lesson progress goals report cards and guardian visibility.",
        ],
    )
    add_heading(doc, "Institution Operations", level=2)
    add_table(
        doc,
        ["Domain", "Included capability"],
        [
            ["People", "Admissions student lifecycle staff memberships batches roles and guardian consent"],
            ["Academics", "Academic years timetable classes attendance exams papers marks results and hall tickets"],
            ["Finance", "Tuition plans installments demand invoices receipts cash verification online orders and reminders"],
            ["Campus", "Notices calendars staff leave substitutions transport hostel passes and visitor records"],
            ["Communication", "Account wide consent communication preferences email and approved WhatsApp workflows"],
        ],
        [1.35, 5.65],
    )
    add_heading(doc, "Seyappaduporul Commercial Model", level=2)
    add_paragraph(
        doc,
        "Revenue can be generated through tutoring fees, recurring institution operations plans, ebook and note commerce, question paper and assessment services, and franchise or managed service agreements. Tuition remains in its dedicated ledger and is consolidated into platform reporting without losing student and invoice detail.",
    )

    page_break(doc)

    add_heading(doc, "Utporul Skills and Careers")
    add_paragraph(
        doc,
        "Utporul is the skill development and creator commerce business. It owns course creation, assessments, credentials, live coding evaluation, and career pathways. Shared authoring primitives allow immersive or tutoring components to be reused, but Utporul remains the product owner for skill programs and professional outcomes.",
    )
    add_heading(doc, "Creation and Assessment", level=2)
    add_bullets(
        doc,
        [
            "Course creation publication collaboration templates packages and versioned restoration.",
            "Quizzes assignments rubrics question banks Assessment Studio and coverage validation.",
            "Games H5P activities mastery evidence adaptive planning and intervention workflows.",
            "Certificates shareable verification and career or internship pathways.",
        ],
    )
    add_heading(doc, "Coding Assessment Isolation", level=2)
    add_paragraph(
        doc,
        "Coding challenges support authored public and hidden tests. Learner submissions enter a durable queue and are leased to a private worker. The worker submits execution to an isolated Judge0 deployment, applies configured resource limits, records retries and results, and communicates with the backend through an internal token. Hidden tests and runner credentials never enter learner responses.",
    )
    add_table(
        doc,
        ["Control", "Purpose"],
        [
            ["Private worker lease", "Prevents learner clients from reaching the execution system"],
            ["Hidden test storage", "Protects assessment integrity"],
            ["Internal service token", "Authenticates worker only endpoints"],
            ["Resource and network limits", "Contains untrusted code execution"],
            ["Retry and status history", "Supports recovery and operational diagnosis"],
        ],
        [2.05, 4.95],
    )
    add_heading(doc, "Utporul Commercial Model", level=2)
    add_paragraph(
        doc,
        "Revenue can be generated through skill course sales, assessment services, premium certificates, creator commerce shares, and career or internship services. Course item allocation provides the basis for instructor earnings and platform margin reporting.",
    )

    page_break(doc)

    add_heading(doc, "Administration Revenue and Support")
    add_heading(doc, "Sasha Control Center", level=2)
    add_paragraph(
        doc,
        "The shared admin console provides portfolio, tenant and entitlement, commercial, per pillar, daily operations, inventory, learning outcomes, revenue, system health, and launch readiness views. Specialized pages retain detailed control over people, courses, enrollments, content libraries, live classes, assessments, certificates, internships, orders, invoices, coupons, memberships, bundles, refunds, and payouts.",
    )
    add_heading(doc, "Commercial Accounting", level=2)
    add_table(
        doc,
        ["Record", "Role in the system", "Integrity rule"],
        [
            ["Offer", "Defines pillar stream price currency tax and entitlement grants", "SKU and vertical classification are explicit"],
            ["Contract", "Binds a tenant to an offer quantity interval and start date", "Customer and commercial terms remain immutable history"],
            ["Invoice", "Creates fiscal numbering subtotal tax discount total and due date", "Number allocation is race safe"],
            ["Ledger event", "Records capture refund or adjustment", "Idempotent source key and immutable amounts"],
            ["Portfolio view", "Aggregates all authoritative revenue sources", "Does not rewrite source transactions"],
        ],
        [1.1, 3.15, 2.75],
    )
    add_heading(doc, "Payments and Fulfillment", level=2)
    add_paragraph(
        doc,
        "Razorpay order values are calculated from server prices. Multi course carts use an immutable compact snapshot of the exact course set and amounts. Verification checks the gateway order, user, amount, signature, and course set before fulfillment. Verify, webhook, and reconciliation paths converge idempotently so a retry cannot duplicate orders, grants, or revenue.",
    )
    add_heading(doc, "Instructor Payouts", level=2)
    add_paragraph(
        doc,
        "Instructor available balance is derived from completed non mock order items for courses owned by that instructor, less pending approved or paid withdrawals. Withdrawal status is one way from pending to approved or rejected and from approved to paid. Administrative transitions create audit and outbox events, and account identifiers are masked in the user interface.",
    )
    add_heading(doc, "Support and View As", level=2)
    add_paragraph(
        doc,
        "Audited view as sessions let administrators reproduce learner or instructor problems without losing the original admin session. Role and tenant checks continue to apply to the impersonated view. This gives support teams a practical diagnostic tool without creating shared passwords.",
    )

    page_break(doc)

    add_heading(doc, "Security Scale and Reliability")
    add_paragraph(
        doc,
        "The release establishes the controls required for a production candidate, but public scale must be proven against the final infrastructure. Capacity depends on the selected database tier, Redis, CDN and object storage, worker counts, meeting infrastructure, media traffic, and code execution limits.",
    )
    add_table(
        doc,
        ["Area", "Implemented control", "Production evidence required"],
        [
            ["Identity", "JWT access and refresh roles two factor admin and session controls", "Key rotation login monitoring and recovery drill"],
            ["Tenancy", "Tenant memberships domains entitlements and scoped services", "Cross tenant penetration and data export tests"],
            ["Money", "Server pricing signatures idempotency immutable ledger and payout states", "Live gateway reconciliation refund and settlement rehearsal"],
            ["Media", "Backend gated video access CDN signing and restricted streaming route", "Bandwidth concurrency and token expiry test"],
            ["Code", "Private runner queue hidden tests internal token and limits", "Sandbox escape network isolation and saturation test"],
            ["Operations", "Audit outbox health views migrations and durable volumes", "Alerting backup restore failover and incident exercise"],
        ],
        [1.05, 3.2, 2.75],
    )
    add_heading(doc, "Scale Strategy", level=2)
    add_bullets(
        doc,
        [
            "Serve static frontend and public assets through CDN and cache policies.",
            "Scale stateless FastAPI workers behind nginx while keeping migrations single runner.",
            "Use managed PostgreSQL with connection pooling replicas for read heavy reporting and tested backups.",
            "Use Redis for transient coordination and separate durable queues where workload warrants it.",
            "Scale streaming meeting and code execution services independently from the business API.",
            "Partition operational dashboards from expensive exports and long running background work.",
            "Set service level objectives and capacity thresholds from measured staging traffic rather than estimates.",
        ],
    )
    add_heading(doc, "Security Operations", level=2)
    add_paragraph(
        doc,
        "Production uses persistent unique application secrets, secure provider credentials, rate limiting, security headers, strict allowed hosts and origins, least privilege service tokens, non root containers, restricted ingress, event logging, and optional security alert email. Secret placeholders and randomly regenerated signing keys are unacceptable in production because they either expose access or invalidate sessions after restart.",
    )

    page_break(doc)

    add_heading(doc, "Production Deployment")
    add_heading(doc, "Required Topology", level=2)
    add_table(
        doc,
        ["Component", "Production role", "Exposure"],
        [
            ["nginx", "TLS termination routing static frontend and API limits", "Public"],
            ["Frontend", "Compiled React application", "Through nginx or CDN"],
            ["Backend", "FastAPI business and admin APIs", "Private behind nginx"],
            ["PostgreSQL", "Authoritative durable data", "Private network only"],
            ["Redis", "Transient cache coordination and locks", "Private network only"],
            ["Streaming", "Server side media extraction and proxy", "Private backend gated"],
            ["Judge0 and worker", "Untrusted coding execution", "Isolated private network"],
            ["Jitsi and Jibri", "Live classes and recordings", "Dedicated secured host"],
            ["Object storage CDN", "Videos ebooks uploads and generated documents", "Signed or policy controlled"],
        ],
        [1.35, 3.4, 2.25],
    )
    add_heading(doc, "Release Sequence", level=2)
    add_numbered(
        doc,
        [
            "Provision PostgreSQL Redis durable storage monitoring and a tested backup destination.",
            "Supply unique persistent platform video live class internal runner database admin Razorpay and SMTP secrets.",
            "Deploy and isolate Judge0 then configure the private worker URL and token.",
            "Take a database backup run Alembic upgrade head and confirm revision 0047.",
            "Build the frontend and start the production Compose stack behind nginx.",
            "Configure DNS and TLS for the LMS API three pillar subdomains and live class host.",
            "Configure payment webhooks sender authentication recording storage CDN monitoring and log retention.",
            "Run staging transaction tenant isolation restore capacity and failure recovery rehearsals before public traffic.",
        ],
    )
    add_heading(doc, "Required Production Values", level=2)
    add_paragraph(
        doc,
        "At minimum the deployment must provide POSTGRES PASSWORD, SECRET KEY, JWT SECRET, VIDEO SECRET, JITSI JWT SECRET, INTERNAL TOKEN, CODE RUNNER TOKEN, JUDGE0 URL, ADMIN PASSWORD, Razorpay key secret and webhook secret, and SMTP host user and password. Bunny Firebase WhatsApp and Sentry values are required only when those capabilities are enabled.",
    )

    page_break(doc)

    add_heading(doc, "Release Evidence")
    add_paragraph(
        doc,
        "The current codebase passed the following software release gates on 13 September 2026. These results establish a reproducible production candidate; they do not substitute for infrastructure specific capacity and disaster recovery evidence.",
    )
    add_table(
        doc,
        ["Gate", "Result"],
        [
            ["Backend regression", "1,879 passed and 4 environment dependent tests skipped"],
            ["Frontend regression", "85 test files and 483 tests passed"],
            ["Frontend quality", "ESLint zero warnings and TypeScript no emit passed"],
            ["Database", "Clean scratch and upgraded seeded preview reached Alembic 0047 head"],
            ["Production build", "59 labs and 143 files checked with 3,483 Vite modules built"],
            ["Deployment definition", "Primary production Compose configuration validated with required values"],
            ["Browser walkthrough", "Admin login portfolio tenants commercial Meiporul operations and payouts loaded against current APIs"],
        ],
        [2.0, 5.0],
    )
    add_heading(doc, "Known Nonblocking Build Warnings", level=2)
    add_paragraph(
        doc,
        "The build reports stale browser compatibility metadata, a mixed static and dynamic auth import, and two lazy chunks above the advisory size threshold. These should be optimized during ongoing performance work but do not represent functional release failures.",
    )
    add_heading(doc, "Go Live Boundary", level=2)
    add_paragraph(
        doc,
        "Public launch is complete only after real DNS and TLS, provider credentials, payment and email delivery, Jitsi and Judge0 isolation, CDN policy, monitoring, database backup restoration, and capacity testing have passed in the target environment. Those actions require production accounts and infrastructure authority and are intentionally not embedded in the source archive.",
    )

    page_break(doc)
    add_heading(doc, "Product Route Map")
    add_table(
        doc,
        ["Purpose", "Route"],
        [
            ["Unified admin control", "/admin/operations"],
            ["Tenant and entitlement control", "/api/v1/platform/tenants"],
            ["Commercial control", "/api/v1/platform/commercial"],
            ["Meiporul deployment operations", "/api/v1/meiporul/operations"],
            ["Utporul coding", "/api/v1/utporul/coding"],
            ["Code worker internal API", "/api/v1/internal/coding"],
            ["Instructor payouts", "/instructor/payouts"],
            ["Admin payout operations", "/admin/payouts"],
            ["Meiporul public experience", "meiporul.sashainfinity.com"],
            ["Seyappaduporul public experience", "seyappaduporul.sashainfinity.com"],
            ["Utporul public experience", "utporul.sashainfinity.com"],
        ],
        [3.0, 4.0],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
